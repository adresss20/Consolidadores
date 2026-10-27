"""
Carga robusta de los archivos de scraping de cada farmacia.

Cada función está diseñada para nunca tronar todo el proceso por culpa de
una sola farmacia: los problemas se acumulan en `stats.errores` /
`stats.warnings` y se devuelven al llamador, quien decide si continúa.
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import openpyxl

from config import FARMACIAS, LONGITUDES_EAN_VALIDAS


@dataclass
class Producto:
    farmacia: str
    sku_origen: str
    ean: Optional[str]          # None si no hay EAN válido
    ean_original: Optional[str]  # valor crudo, aunque no sea válido (para trazabilidad)
    nombre: str
    marca: str
    precio_lista: Optional[float]
    precio_final: Optional[float]


@dataclass
class ResultadoCarga:
    farmacia: str
    productos: list = field(default_factory=list)
    archivo_usado: Optional[Path] = None
    total_filas: int = 0
    ean_validos: int = 0
    ean_invalidos: int = 0
    duplicados_ean: int = 0
    precios_invalidos: int = 0
    errores: list = field(default_factory=list)   # impiden usar la farmacia
    warnings: list = field(default_factory=list)   # no impiden usarla
    errores_scraping_fuente: Optional[int] = None  # ej. hoja "Errores" de San Pablo


def _quitar_acentos(txt: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", txt) if unicodedata.category(c) != "Mn"
    )


def _localizar_archivo(carpeta: Path, resultado: ResultadoCarga) -> Optional[Path]:
    if not carpeta.exists():
        resultado.errores.append(
            f"No existe la carpeta de entrada '{carpeta}'. Créala y coloca ahí el .xlsx de esta farmacia."
        )
        return None

    candidatos = sorted(carpeta.glob("*.xlsx")) + sorted(carpeta.glob("*.xlsm"))
    # ignorar archivos temporales de Excel (~$archivo.xlsx)
    candidatos = [c for c in candidatos if not c.name.startswith("~$")]

    if not candidatos:
        resultado.errores.append(
            f"No se encontró ningún archivo .xlsx en '{carpeta}'. "
            "Coloca ahí el archivo de scraping más reciente de esta farmacia."
        )
        return None

    if len(candidatos) > 1:
        candidatos.sort(key=lambda p: p.stat().st_mtime, reverse=True)
        elegido = candidatos[0]
        ignorados = ", ".join(c.name for c in candidatos[1:])
        resultado.warnings.append(
            f"Se encontraron {len(candidatos)} archivos en '{carpeta}'. "
            f"Se usó el más reciente ('{elegido.name}') e se ignoraron: {ignorados}."
        )
        return elegido

    return candidatos[0]


def limpiar_ean(valor) -> tuple[Optional[str], Optional[str]]:
    """Devuelve (ean_valido_o_None, valor_original_como_texto_o_None)."""
    if valor is None:
        return None, None
    texto = str(valor).strip()
    if not texto or texto.lower() in {"none", "nan", "n/a", "-"}:
        return None, None
    original = texto
    # Excel a veces guarda EAN como float (7501033953064.0)
    if texto.endswith(".0"):
        texto = texto[:-2]
    texto = re.sub(r"\D", "", texto)  # dejar solo dígitos
    if len(texto) in LONGITUDES_EAN_VALIDAS:
        return texto, original
    return None, original


def limpiar_precio(valor) -> Optional[float]:
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        num = float(valor)
    else:
        texto = str(valor).strip().replace("$", "").replace(",", "")
        if not texto:
            return None
        try:
            num = float(texto)
        except ValueError:
            return None
    if num <= 0:
        return None
    return round(num, 2)


def normalizar_nombre(nombre: str) -> str:
    if not nombre:
        return ""
    txt = _quitar_acentos(nombre.lower())
    txt = re.sub(r"[^a-z0-9\s]", " ", txt)
    txt = re.sub(r"\s+", " ", txt).strip()
    return txt


def normalizar_marca(marca: str) -> str:
    if not marca:
        return ""
    return _quitar_acentos(str(marca).lower()).strip()


def cargar_farmacia(farmacia_key: str, entradas_dir: Path) -> ResultadoCarga:
    cfg = FARMACIAS[farmacia_key]
    resultado = ResultadoCarga(farmacia=farmacia_key)

    carpeta = entradas_dir / farmacia_key
    archivo = _localizar_archivo(carpeta, resultado)
    if archivo is None:
        return resultado
    resultado.archivo_usado = archivo

    try:
        wb = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - queremos capturar cualquier fallo de lectura
        resultado.errores.append(f"No se pudo abrir '{archivo.name}': {exc}")
        return resultado

    try:
        hoja_nombre = cfg["hoja"]
        if hoja_nombre not in wb.sheetnames:
            resultado.errores.append(
                f"El archivo '{archivo.name}' no tiene una hoja llamada '{hoja_nombre}'. "
                f"Hojas disponibles: {wb.sheetnames}"
            )
            return resultado
        ws = wb[hoja_nombre]

        filas = ws.iter_rows(values_only=True)
        for _ in range(cfg["header_row"] + 1):
            try:
                encabezados = next(filas)
            except StopIteration:
                resultado.errores.append(
                    f"La hoja '{hoja_nombre}' de '{archivo.name}' no tiene suficientes filas para el encabezado configurado."
                )
                return resultado

        encabezados = [str(h).strip() if h is not None else "" for h in encabezados]
        indice = {h: i for i, h in enumerate(encabezados)}

        columnas_requeridas = ["ean", "nombre", "precio_lista", "precio_final"]
        faltantes = [
            cfg[campo] for campo in columnas_requeridas
            if cfg[campo] is not None and cfg[campo] not in indice
        ]
        if faltantes:
            resultado.errores.append(
                f"Faltan columnas esperadas en '{archivo.name}' (hoja '{hoja_nombre}'): {faltantes}. "
                f"Columnas encontradas: {encabezados}"
            )
            return resultado

        idx_ean = indice.get(cfg["ean"])
        idx_nombre = indice[cfg["nombre"]]
        idx_marca = indice.get(cfg["marca"]) if cfg.get("marca") else None
        idx_precio_lista = indice[cfg["precio_lista"]]
        idx_precio_final = indice[cfg["precio_final"]]
        idx_sku = indice.get(cfg.get("sku_origen"))

        eans_vistos: dict[str, int] = {}  # ean -> índice en resultado.productos (para dedupe)

        for fila in filas:
            if fila is None or all(v is None for v in fila):
                continue
            resultado.total_filas += 1

            nombre = fila[idx_nombre] if idx_nombre < len(fila) else None
            if nombre is None or str(nombre).strip() == "":
                continue  # fila sin nombre no es un producto usable
            nombre = str(nombre).strip()

            marca = ""
            if idx_marca is not None and idx_marca < len(fila) and fila[idx_marca] is not None:
                marca = str(fila[idx_marca]).strip()

            sku_origen = ""
            if idx_sku is not None and idx_sku < len(fila) and fila[idx_sku] is not None:
                sku_origen = str(fila[idx_sku]).strip()

            ean_valido, ean_original = (None, None)
            if idx_ean is not None:
                ean_valido, ean_original = limpiar_ean(fila[idx_ean] if idx_ean < len(fila) else None)

            if ean_valido is not None:
                resultado.ean_validos += 1
            else:
                # cuenta tanto celdas vacías como valores con formato inválido
                resultado.ean_invalidos += 1

            precio_lista = limpiar_precio(fila[idx_precio_lista] if idx_precio_lista < len(fila) else None)
            precio_final = limpiar_precio(fila[idx_precio_final] if idx_precio_final < len(fila) else None)
            if precio_lista is None and precio_final is None:
                resultado.precios_invalidos += 1
            # si falta el precio final pero hay precio de lista, se usa el de lista como final
            if precio_final is None:
                precio_final = precio_lista

            producto = Producto(
                farmacia=farmacia_key,
                sku_origen=sku_origen,
                ean=ean_valido,
                ean_original=ean_original,
                nombre=nombre,
                marca=marca,
                precio_lista=precio_lista,
                precio_final=precio_final,
            )

            if ean_valido is not None:
                if ean_valido in eans_vistos:
                    resultado.duplicados_ean += 1
                    idx_existente = eans_vistos[ean_valido]
                    existente = resultado.productos[idx_existente]
                    precio_existente = existente.precio_final if existente.precio_final is not None else float("inf")
                    precio_nuevo = producto.precio_final if producto.precio_final is not None else float("inf")
                    if precio_nuevo < precio_existente:
                        resultado.productos[idx_existente] = producto
                    continue
                eans_vistos[ean_valido] = len(resultado.productos)

            resultado.productos.append(producto)

        # hoja opcional de errores de scraping de la propia fuente (ej. San Pablo)
        hoja_errores = cfg.get("hoja_errores")
        if hoja_errores and hoja_errores in wb.sheetnames:
            try:
                ws_err = wb[hoja_errores]
                primera = next(ws_err.iter_rows(min_row=1, max_row=1, values_only=True), None)
                if primera and primera[0]:
                    match = re.search(r"(\d+)", str(primera[0]))
                    if match:
                        resultado.errores_scraping_fuente = int(match.group(1))
            except Exception:  # noqa: BLE001 - esto es solo informativo, nunca debe tronar la carga
                pass

    finally:
        wb.close()

    return resultado
