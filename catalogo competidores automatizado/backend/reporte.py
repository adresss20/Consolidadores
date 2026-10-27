"""
Escribe el .xlsx final del catálogo consolidado con 3 hojas:
  - Catálogo Consolidado
  - SKUs sin identificar
  - Resumen del proceso
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import FARMACIAS
from consolidar import ORDEN_FARMACIAS, FilaConsolidada
from io_utils import ResultadoCarga
from matching import SinIdentificar

FMT_MONEDA = '"$"#,##0.00'

ENCABEZADO_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
ENCABEZADO_FONT = Font(bold=True, color="FFFFFF")
BARATO_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")


def _escribir_encabezado(ws, encabezados: list[str]):
    for col, titulo in enumerate(encabezados, start=1):
        celda = ws.cell(row=1, column=col, value=titulo)
        celda.font = ENCABEZADO_FONT
        celda.fill = ENCABEZADO_FILL
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def _autoajustar_columnas(ws, anchos: list[int]):
    for col, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(col)].width = ancho


def _hoja_catalogo_consolidado(wb: Workbook, filas: list[FilaConsolidada]):
    ws = wb.active
    ws.title = "Catálogo Consolidado"

    encabezados = ["EAN"]
    for farmacia in ORDEN_FARMACIAS:
        visible = FARMACIAS[farmacia]["nombre_visible"]
        encabezados += [f"Nombre {visible}", f"Precio Lista {visible}", f"Precio Final {visible}"]
    encabezados += ["Farmacia Más Barata", "Precio Más Barato", "Coincidencias asumidas por nombre"]
    _escribir_encabezado(ws, encabezados)

    for fila_idx, fila in enumerate(filas, start=2):
        col = 1
        ws.cell(row=fila_idx, column=col, value=fila.ean)
        col += 1
        for farmacia in ORDEN_FARMACIAS:
            producto = fila.por_farmacia.get(farmacia)
            ws.cell(row=fila_idx, column=col, value=producto.nombre if producto else None)
            ws.cell(row=fila_idx, column=col + 1, value=producto.precio_lista if producto else None)
            ws.cell(row=fila_idx, column=col + 1).number_format = FMT_MONEDA
            ws.cell(row=fila_idx, column=col + 2, value=producto.precio_final if producto else None)
            ws.cell(row=fila_idx, column=col + 2).number_format = FMT_MONEDA
            col += 3

        farmacia_barata, precio_barato = fila.farmacia_mas_barata()
        celda_farmacia = ws.cell(row=fila_idx, column=col, value=farmacia_barata)
        celda_farmacia.fill = BARATO_FILL
        celda_precio = ws.cell(row=fila_idx, column=col + 1, value=precio_barato)
        celda_precio.number_format = FMT_MONEDA
        celda_precio.fill = BARATO_FILL
        ws.cell(row=fila_idx, column=col + 2, value="; ".join(fila.metodos_asumidos) or None)

    anchos = [16]
    for _ in ORDEN_FARMACIAS:
        anchos += [38, 14, 14]
    anchos += [22, 16, 40]
    _autoajustar_columnas(ws, anchos)
    ws.auto_filter.ref = ws.dimensions


def _hoja_sin_identificar(wb: Workbook, sin_identificar: list[SinIdentificar]):
    ws = wb.create_sheet("SKUs sin identificar")
    encabezados = [
        "Farmacia", "SKU Origen", "EAN detectado (inválido)", "Nombre", "Marca",
        "Precio Lista", "Precio Final", "Motivo", "Candidatos considerados",
    ]
    _escribir_encabezado(ws, encabezados)

    for fila_idx, item in enumerate(sin_identificar, start=2):
        p = item.producto
        ws.cell(row=fila_idx, column=1, value=FARMACIAS[p.farmacia]["nombre_visible"])
        ws.cell(row=fila_idx, column=2, value=p.sku_origen)
        ws.cell(row=fila_idx, column=3, value=p.ean_original)
        ws.cell(row=fila_idx, column=4, value=p.nombre)
        ws.cell(row=fila_idx, column=5, value=p.marca)
        c = ws.cell(row=fila_idx, column=6, value=p.precio_lista)
        c.number_format = FMT_MONEDA
        c = ws.cell(row=fila_idx, column=7, value=p.precio_final)
        c.number_format = FMT_MONEDA
        ws.cell(row=fila_idx, column=8, value=item.motivo)
        ws.cell(row=fila_idx, column=9, value=item.candidatos)

    _autoajustar_columnas(ws, [16, 14, 18, 45, 20, 14, 14, 32, 60])
    ws.auto_filter.ref = ws.dimensions


def _hoja_resumen(wb: Workbook, resultados_carga: dict[str, ResultadoCarga], n_consolidadas: int,
                   n_sin_identificar: int, n_match_por_nombre: int):
    ws = wb.create_sheet("Resumen del proceso")
    ws.cell(row=1, column=1, value="RESUMEN DE LA CORRIDA").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value="Fecha de generación")
    ws.cell(row=2, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M"))

    ws.cell(row=4, column=1, value="Productos en el catálogo consolidado").font = Font(bold=True)
    ws.cell(row=4, column=2, value=n_consolidadas)
    ws.cell(row=5, column=1, value="  de los cuales emparejados por nombre (no EAN real)")
    ws.cell(row=5, column=2, value=n_match_por_nombre)
    ws.cell(row=6, column=1, value="Filas en 'SKUs sin identificar'")
    ws.cell(row=6, column=2, value=n_sin_identificar)

    fila = 8
    encabezados = [
        "Farmacia", "Archivo usado", "Filas leídas", "EAN válidos", "EAN inválidos/faltantes",
        "Duplicados de EAN", "Precios inválidos", "Errores de scraping (fuente)", "Advertencias", "Errores (excluida)",
    ]
    _escribir_encabezado_en(ws, fila, encabezados)
    fila += 1
    for farmacia_key in ORDEN_FARMACIAS:
        r = resultados_carga.get(farmacia_key)
        if r is None:
            continue
        ws.cell(row=fila, column=1, value=FARMACIAS[farmacia_key]["nombre_visible"])
        ws.cell(row=fila, column=2, value=r.archivo_usado.name if r.archivo_usado else "(no encontrado)")
        ws.cell(row=fila, column=3, value=r.total_filas)
        ws.cell(row=fila, column=4, value=r.ean_validos)
        ws.cell(row=fila, column=5, value=r.ean_invalidos)
        ws.cell(row=fila, column=6, value=r.duplicados_ean)
        ws.cell(row=fila, column=7, value=r.precios_invalidos)
        ws.cell(row=fila, column=8, value=r.errores_scraping_fuente)
        ws.cell(row=fila, column=9, value=" | ".join(r.warnings) or None)
        ws.cell(row=fila, column=10, value=" | ".join(r.errores) or None)
        fila += 1

    _autoajustar_columnas(ws, [16, 26, 12, 12, 18, 14, 14, 20, 50, 50])


def _escribir_encabezado_en(ws, fila: int, encabezados: list[str]):
    for col, titulo in enumerate(encabezados, start=1):
        celda = ws.cell(row=fila, column=col, value=titulo)
        celda.font = ENCABEZADO_FONT
        celda.fill = ENCABEZADO_FILL


def generar_reporte(
    filas: list[FilaConsolidada],
    sin_identificar: list[SinIdentificar],
    resultados_carga: dict[str, ResultadoCarga],
    salida_dir: Path,
) -> Path:
    wb = Workbook()
    _hoja_catalogo_consolidado(wb, filas)
    _hoja_sin_identificar(wb, sin_identificar)

    n_match_por_nombre = sum(1 for f in filas if f.metodos_asumidos)
    _hoja_resumen(wb, resultados_carga, len(filas), len(sin_identificar), n_match_por_nombre)

    salida_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    ruta_salida = salida_dir / f"Catalogo_Consolidado_{timestamp}.xlsx"
    wb.save(ruta_salida)
    return ruta_salida
