"""
exporter.py
-----------
Genera el archivo Excel de consolidado con formato visual limpio:
  - Header oscuro con texto blanco
  - Columnas de precio con formato moneda y 2 decimales
  - Proveedor más barato resaltado en verde
  - Segundo proveedor resaltado en amarillo
  - Filas alternadas en gris claro para legibilidad
  - Primera fila congelada + auto-filter
  - Anchos de columna automáticos
"""

import pandas as pd
import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill,
    Font,
    Alignment,
    Border,
    Side,
    numbers,
)
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# ── Paleta de colores ──────────────────────────────────────────────────────────
COLOR_HEADER_BG      = "1F3864"   # Azul oscuro
COLOR_HEADER_FG      = "FFFFFF"
COLOR_ROW_ALT        = "F2F2F2"   # Gris muy claro (filas pares)
COLOR_ROW_WHITE      = "FFFFFF"
COLOR_CHEAPEST       = "C6EFCE"   # Verde claro
COLOR_CHEAPEST_FONT  = "276221"
COLOR_SECOND         = "FFEB9C"   # Amarillo claro
COLOR_SECOND_FONT    = "9C5700"
COLOR_PRICE_COL_BG   = "EBF1DE"   # Verde muy pálido para columnas de precio
COLOR_BORDER         = "CCCCCC"

# ── Fills pre-construidos ──────────────────────────────────────────────────────
FILL_HEADER   = PatternFill("solid", fgColor=COLOR_HEADER_BG)
FILL_ALT      = PatternFill("solid", fgColor=COLOR_ROW_ALT)
FILL_WHITE    = PatternFill("solid", fgColor=COLOR_ROW_WHITE)
FILL_CHEAPEST = PatternFill("solid", fgColor=COLOR_CHEAPEST)
FILL_SECOND   = PatternFill("solid", fgColor=COLOR_SECOND)

FONT_HEADER   = Font(bold=True, color=COLOR_HEADER_FG, name="Arial", size=10)
FONT_DATA     = Font(name="Arial", size=9)
FONT_CHEAPEST = Font(bold=True, color=COLOR_CHEAPEST_FONT, name="Arial", size=9)
FONT_SECOND   = Font(bold=True, color=COLOR_SECOND_FONT, name="Arial", size=9)

BORDER_THIN = Border(
    bottom=Side(style="thin", color=COLOR_BORDER),
)

ALIGN_CENTER  = Alignment(horizontal="center", vertical="center", wrap_text=False)
ALIGN_LEFT    = Alignment(horizontal="left",   vertical="center")
ALIGN_RIGHT   = Alignment(horizontal="right",  vertical="center")
ALIGN_HEADER  = Alignment(horizontal="center", vertical="center", wrap_text=True)

FMT_PRECIO    = '#,##0.00'
FMT_INT       = '0'


def export_to_excel(
    df: pd.DataFrame,
    output_path: Path,
    supplier_names: list,
) -> None:
    """
    Escribe el DataFrame consolidado en output_path con formato profesional.
    """
    # ── 1. Definir columnas de salida y etiquetas ──────────────────────────────
    info_cols = [
        ("ean",          "SKU / EAN"),
        ("descripcion",  "DESCRIPCIÓN"),
        ("laboratorio",  "LABORATORIO"),
        ("familia",      "FAMILIA"),
        ("generico",     "GENÉRICO"),
        ("sustancia",    "SUSTANCIA ACTIVA"),
        ("presentacion", "PRESENTACIÓN"),
    ]
    price_col_defs = [
        (f"precio_{name}", name.upper())
        for name in supplier_names
        if f"precio_{name}" in df.columns
    ]
    analysis_cols = [
        ("PROVEEDOR_MAS_BARATO", "PROV. MÁS BARATO"),
        ("PRECIO_MAS_BARATO",    "PRECIO MÁS BARATO"),
        ("SEGUNDO_MAS_BARATO",   "2DO MÁS BARATO"),
        ("PRECIO_SEGUNDO",       "PRECIO 2DO"),
        ("AHORRO_VS_SEGUNDO",    "AHORRO vs 2DO"),
        ("NUM_PROVEEDORES",      "# PROVEEDORES"),
    ]

    all_col_defs = info_cols + price_col_defs + analysis_cols
    out_df_cols   = [c for c, _ in all_col_defs if c in df.columns]
    header_labels = {c: lbl for c, lbl in all_col_defs if c in df.columns}

    df_out = df[out_df_cols].copy()
    df_out = df_out.rename(columns=header_labels)

    # ── 2. Identificar conjuntos de columnas para formato ──────────────────────
    price_labels       = {lbl for c, lbl in price_col_defs}
    analysis_labels    = {lbl for c, lbl in analysis_cols}
    cheapest_label     = header_labels.get("PROVEEDOR_MAS_BARATO")
    precio_1_label     = header_labels.get("PRECIO_MAS_BARATO")
    segundo_label      = header_labels.get("SEGUNDO_MAS_BARATO")
    precio_2_label     = header_labels.get("PRECIO_SEGUNDO")
    ahorro_label       = header_labels.get("AHORRO_VS_SEGUNDO")
    num_prov_label     = header_labels.get("NUM_PROVEEDORES")

    price_money_labels = (
        price_labels
        | {precio_1_label, precio_2_label, ahorro_label}
        - {None}
    )

    # ── 3. Crear workbook ──────────────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado"

    headers = list(df_out.columns)
    n_cols  = len(headers)
    n_rows  = len(df_out)

    # ── 4. Escribir header ─────────────────────────────────────────────────────
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill      = FILL_HEADER
        cell.font      = FONT_HEADER
        cell.alignment = ALIGN_HEADER

    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}1"

    # ── 5. Escribir datos ──────────────────────────────────────────────────────
    logger.info(f"Escribiendo {n_rows:,} filas en Excel...")

    for row_idx, row_data in enumerate(df_out.itertuples(index=False), start=2):
        is_even     = (row_idx % 2 == 0)
        base_fill   = FILL_ALT if is_even else FILL_WHITE
        row_dict    = dict(zip(headers, row_data))

        # Determinar qué proveedor es el más barato y el segundo
        prov_1  = row_dict.get(cheapest_label)
        prov_2  = row_dict.get(segundo_label)

        for col_idx, header in enumerate(headers, start=1):
            value = row_dict[header]

            # Convertir pd.NA / NaN a None para que Excel muestre celda vacía
            if pd.isna(value) if not isinstance(value, str) else False:
                value = None

            cell = ws.cell(row=row_idx, column=col_idx, value=value)

            # Formato de número
            if header in price_money_labels and value is not None:
                cell.number_format = FMT_PRECIO
                cell.alignment     = ALIGN_RIGHT
            elif header == num_prov_label and value is not None:
                cell.number_format = FMT_INT
                cell.alignment     = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT

            cell.font   = FONT_DATA
            cell.border = BORDER_THIN

            # Resaltado proveedor más barato
            if prov_1 and header == supplier_name_to_label(prov_1, price_col_defs):
                cell.fill = FILL_CHEAPEST
                cell.font = FONT_CHEAPEST
            # Resaltado segundo más barato
            elif prov_2 and header == supplier_name_to_label(prov_2, price_col_defs):
                cell.fill = FILL_SECOND
                cell.font = FONT_SECOND
            else:
                cell.fill = base_fill

    # ── 6. Anchos de columna ───────────────────────────────────────────────────
    _auto_column_widths(ws, headers, df_out)

    # ── 7. Guardar ─────────────────────────────────────────────────────────────
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    logger.info(f"Archivo guardado: {output_path}")


# ── Helpers ────────────────────────────────────────────────────────────────────

def supplier_name_to_label(supplier_name: str, price_col_defs: list) -> str | None:
    """
    Dado un nombre de proveedor (e.g. 'Adsa'), retorna la etiqueta de columna
    de precio correspondiente (e.g. 'ADSA').
    """
    lookup = {name: lbl for col, lbl in price_col_defs
              for name in [col.replace("precio_", "")]}
    return lookup.get(supplier_name)


def _auto_column_widths(ws, headers: list, df: pd.DataFrame) -> None:
    """Ajusta el ancho de cada columna según el contenido."""
    MIN_WIDTH = 10
    MAX_WIDTH = 45

    for col_idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)

        # Ancho del header (con salto de línea virtual)
        header_width = max(len(part) for part in header.split()) + 2

        # Ancho máximo en los datos (muestra de hasta 500 filas para velocidad)
        try:
            sample = df[header].dropna().astype(str).head(500)
            data_width = sample.str.len().max() if not sample.empty else 0
        except Exception:
            data_width = 0

        width = max(header_width, data_width if data_width else 0, MIN_WIDTH)
        width = min(width, MAX_WIDTH)

        ws.column_dimensions[col_letter].width = width
